import sys, collections
sys.path.insert(0, '.')
from openpyxl import load_workbook
from parse_br import load
from brands import brand

IVA = 1.21
DTO = 0.55          # descuento de Hendi sobre su Net price
FEE = 0.1545        # comision de referencia estimada sobre PVP bruto

def num(x):
    if x is None: return None
    if isinstance(x, (int, float)): return float(x)
    s = str(x).strip().replace('€', '').replace('\xa0', '')
    if not s: return None
    # "399,04" europeo vs "428.35" anglosajon
    if ',' in s and '.' in s: s = s.replace('.', '').replace(',', '.')
    elif ',' in s: s = s.replace(',', '.')
    try: return float(s)
    except ValueError: return None

# ---- hoja de costes, indexada por MPN
H = {}
ws = load_workbook('costes/hendi_raw/hendi.xlsx', data_only=True, read_only=True)['Hoja1']
for r in ws.iter_rows(min_row=2, values_only=True):
    mpn = (str(r[3]).strip() if r[3] is not None else '')
    if not mpn: continue
    H[mpn] = dict(gtin=r[0], neto_hendi=num(r[4]), bb_ganada=(str(r[5] or '').strip().lower() == 'si'),
                  bb_precio=num(r[6]))

# ---- ventas Hendi ene-sep 2026
V = collections.defaultdict(lambda: dict(u=0., v=0., s=0., b=0., t=''))
for i in range(1, 10):
    for row in load(f'mensual/2026-{i:02d}.csv'):
        if brand(row) != 'Hendi': continue
        d = V[row['sku']]
        d['u'] += row['units']; d['v'] += row['sales']
        d['s'] += row['sessions']; d['b'] += row['bb'] * row['sessions']
        if not d['t']: d['t'] = row['title'][:40]

def mpn_de(sku):
    """H470411 -> 470411. Tolera sufijos tipo _fba."""
    s = sku.upper().lstrip('H').split('_')[0]
    return s if s in H else None

F = []
for sku, d in V.items():
    if not d['u']: continue
    m = mpn_de(sku)
    if m is None or not H[m]['neto_hendi']: continue
    h = H[m]
    coste = h['neto_hendi'] * (1 - DTO)
    pvp = d['v'] / d['u']                  # PVP bruto medio realizado
    neto = pvp / IVA
    margen = neto - pvp * FEE - coste
    F.append(dict(sku=sku, mpn=m, t=d['t'], uds=d['u'], venta=d['v'], pvp=pvp,
                  neto_hendi=h['neto_hendi'], coste=coste, margen=margen,
                  pct_neto=100 * margen / neto, equilibrio=100 * margen / pvp,
                  contrib=margen * d['u'], bb=d['b'] / d['s'] if d['s'] else 0,
                  bb_comp=h['bb_precio'], bb_ganada=h['bb_ganada']))
F.sort(key=lambda r: -r['venta'])

tv = sum(d['v'] for d in V.values())
cub = sum(r['venta'] for r in F)
print(f"Hoja: {len(H)} referencias | ventas Hendi: {len(V)} SKU, {tv:,.0f} €")
print(f"Emparejadas: {len(F)} SKU, {cub:,.0f} € ({cub/tv:.1%} de la venta Hendi)")
M = sum(r['margen'] * r['uds'] for r in F); G = sum(r['pvp'] * r['uds'] for r in F)
N = sum(r['pvp'] / IVA * r['uds'] for r in F)
print(f"\nMARGEN AGREGADO (dto {DTO:.0%}, fee {FEE:.2%})")
print(f"  venta bruta {G:,.0f} € | margen {M:,.0f} €")
print(f"  margen s/neto      : {100*M/N:5.1f}%")
print(f"  ACOS DE EQUILIBRIO : {100*M/G:5.1f}%")

print(f"\n{'SKU':12s}{'venta':>9s}{'uds':>5s}{'PVP':>8s}{'coste':>8s}{'mrg/ud':>8s}{'%neto':>7s}{'equil':>7s}{'BB':>7s}  producto")
for r in F[:25]:
    print(f"{r['sku'][:12]:12s}{r['venta']:9,.0f}{r['uds']:5.0f}{r['pvp']:8.1f}{r['coste']:8.1f}"
          f"{r['margen']:8.2f}{r['pct_neto']:6.1f}%{r['equilibrio']:6.1f}%{r['bb']:6.1f}%  {r['t']}")
neg = [r for r in F if r['margen'] <= 0]
print(f"\nSKU CON MARGEN NEGATIVO: {len(neg)}  |  venta {sum(r['venta'] for r in neg):,.0f} €  |  perdida {sum(r['contrib'] for r in neg):,.0f} €")
for r in sorted(neg, key=lambda r: r['contrib'])[:10]:
    print(f"  {r['contrib']:9,.0f} €  {r['sku'][:12]:12s} PVP {r['pvp']:7.1f} coste {r['coste']:7.1f} uds {r['uds']:4.0f}  {r['t']}")
# precio del competidor en buybox
pc = [r for r in F if r['bb_comp'] and r['bb'] < 90 and r['venta'] > 1000]
print(f"\nDONDE EL COMPETIDOR GANA LA BUYBOX ({len(pc)} SKU con venta >1.000 €)")
print(f"{'SKU':12s}{'BB nuestra':>11s}{'PVP nuestro':>12s}{'PVP comp.':>11s}{'dif':>8s}  producto")
for r in sorted(pc, key=lambda r: -r['venta'])[:12]:
    print(f"{r['sku'][:12]:12s}{r['bb']:10.1f}%{r['pvp']:12.1f}{r['bb_comp']:11.1f}{(r['pvp']-r['bb_comp'])/r['bb_comp']*100:+7.1f}%  {r['t']}")
