"""Parsea la tabla de costes de VINFER (una hoja por marca) y la recalcula.
   Todos los costes vienen en BRUTO: se netean dividiendo entre 1,21."""
import openpyxl, collections, json, re, warnings
warnings.filterwarnings('ignore')

IVA = 1.21
XL = 'costes/vinfer_raw/vinfer.xlsx'

def celdas(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]

def cabecera(rows):
    for i, r in enumerate(rows[:6]):
        txt = [str(c).strip().upper() if c else '' for c in r]
        if 'EAN' in txt and any('COSTE' in t for t in txt):
            return i, txt
    return None, None

def num(v):
    if isinstance(v, (int, float)): return float(v)
    return None

def leer():
    wb = openpyxl.load_workbook(XL, read_only=True, data_only=True)
    out = []
    for hoja in wb.sheetnames:
        rows = celdas(wb[hoja])
        hi, hdr = cabecera(rows)
        if hi is None: continue
        col = {}
        for j, t in enumerate(hdr):
            if t.startswith('EAN'): col['ean'] = j
            elif t in ('TÍTULO', 'TITULO'): col['title'] = j
            elif t == 'SKU': col['sku'] = j
            elif t == 'CDAD': col['cdad'] = j
            elif t.startswith('REF'): col['ref'] = j
            elif t.startswith('COSTE') and 'FINAL' not in t: col['coste'] = j
            elif 'AMZON' in t: col['fee'] = j
            elif t.startswith('ENV'): col['envio'] = j
            elif t == 'IVA': col['iva'] = j
            elif 'COSTE FINAL' in t: col['cf'] = j
            elif t == 'MARGEN': col['margen'] = j
            elif t == 'PVP': col['pvp'] = j
            elif '%MARGEN' in t.replace(' ', ''): col['pct'] = j
            elif 'IVA COMPRA' in t: col['ivac'] = j
        grupo = ''
        for r in rows[hi+1:]:
            if not any(r): continue
            ean = str(r[col['ean']]).strip() if col.get('ean') is not None and r[col['ean']] else ''
            coste = num(r[col['coste']]) if col.get('coste') is not None else None
            if coste is None:
                t0 = str(r[0] or '').strip()
                if t0 and not t0.isdigit(): grupo = t0[:40]
                continue
            g = lambda k: num(r[col[k]]) if col.get(k) is not None else None
            s = lambda k: (str(r[col[k]]).replace('\n', ' ').strip()
                           if col.get(k) is not None and r[col[k]] else '')
            out.append(dict(marca=hoja, grupo=grupo, ean=ean, ref=s('ref'), sku_hoja=s('sku'),
                            title=s('title'), cdad=s('cdad'),
                            coste_bruto=coste, fee=g('fee'), envio=g('envio'),
                            iva_hoja=g('iva'), cf_hoja=g('cf'), margen_hoja=g('margen'),
                            pvp=g('pvp'), pct_hoja=g('pct'), iva_compra=g('ivac'),
                            tiene_ivac=col.get('ivac') is not None))
    return out

F = leer()
for f in F:
    p = f['pvp'] or 0
    f['tasa']      = (f['fee']/p) if p and f['fee'] is not None else None
    f['iva_ok']    = p*0.21/1.21
    f['coste_neto']= f['coste_bruto']/IVA
    f['neto']      = p - f['iva_ok']
    f['margen_ok'] = f['neto'] - (f['fee'] or 0) - (f['envio'] or 0) - f['coste_neto']
    f['pct_ok']    = 100*f['margen_ok']/f['neto'] if f['neto'] else 0
    f['pct_sobre_pvp'] = 100*f['margen_ok']/p if p else 0

if __name__ == '__main__':
    print(f"{len(F)} referencias en {len({f['marca'] for f in F})} marcas\n")
    print(f"{'Marca':<14}{'Refs':>6}{'IVA COMPRA':>12}{'IVA hoja OK':>13}{'Tasa comision':>16}")
    print('-'*62)
    for m in dict.fromkeys(f['marca'] for f in F):
        g = [f for f in F if f['marca'] == m]
        ivac = 'si' if g[0]['tiene_ivac'] else 'NO'
        bien = sum(1 for f in g if abs((f['iva_hoja'] or 0) - f['iva_ok']) < 0.01)
        tasas = sorted({round(f['tasa'], 4) for f in g if f['tasa']})
        print(f"{m:<14}{len(g):>6}{ivac:>12}{f'{bien}/{len(g)}':>13}"
              f"{', '.join(f'{100*t:.1f}%' for t in tasas[:4]):>16}")

    print("\n\nPROBLEMA 1 — el IVA de venta esta fijado a mano en 4 hojas\n")
    mal = [f for f in F if abs((f['iva_hoja'] or 0) - f['iva_ok']) > 0.01]
    print(f"  {len(mal)} de {len(F)} referencias ({100*len(mal)/len(F):.0f}%)\n")
    print(f"{'Marca':<13}{'PVP':>7}{'IVA hoja':>10}{'IVA correcto':>14}{'Error':>8}{'efecto margen':>15}")
    vis = set()
    for f in sorted(mal, key=lambda f: -abs((f['iva_hoja'] or 0)-f['iva_ok'])):
        k = (f['marca'], f['pvp'])
        if k in vis: continue
        vis.add(k)
        d = (f['iva_hoja'] or 0) - f['iva_ok']
        print(f"{f['marca']:<13}{f['pvp']:>7.2f}{f['iva_hoja']:>10.4f}{f['iva_ok']:>14.4f}"
              f"{d:>+8.3f}{('infravalora' if d>0 else 'SOBREVALORA')+f' {abs(d):.2f}':>15}")
        if len(vis) > 11: break

    print("\n\nPROBLEMA 2 — el IVA soportado solo se recupera en 2 de las 5 hojas\n")
    for m in dict.fromkeys(f['marca'] for f in F):
        g = [f for f in F if f['marca'] == m]
        perd = sum(f['coste_bruto'] - f['coste_neto'] for f in g)
        print(f"  {m:<14}{'recupera IVA compra' if g[0]['tiene_ivac'] else 'NO lo recupera':<24}"
              f"{len(g):>4} refs   {perd:>7.2f} EUR/ud sumados que se dejan de contar"
              if not g[0]['tiene_ivac'] else
              f"  {m:<14}{'recupera IVA compra':<24}{len(g):>4} refs")

    json.dump(F, open('vinfer_costes.json', 'w'), ensure_ascii=False)
    print(f"\n-> vinfer_costes.json")
