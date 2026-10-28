"""Canal Makro/Metro, con los dos anos separados.
2025 es ano natural completo; 2026 llega al 29/10, asi que la comparacion
honesta es por EUR/mes, no por total."""
import json, collections, re
from openpyxl import load_workbook

FICHEROS = {'2025': ('makro_raw/makro_2025.xlsx', 12.0),
            '2026': ('makro_raw/makro_2026.xlsx', 10.0)}
PLAT = {'makro.es': 'Makro España', 'makro.pt': 'Makro Portugal',
        'metro.it': 'Metro Italia', 'makro.nl': 'Makro Países Bajos'}

def num(x):
    if x is None: return 0.0
    if isinstance(x, (int, float)): return float(x)
    try: return float(str(x).replace('€', '').strip())
    except ValueError: return 0.0

def marca(r):
    mm = str(r[10] or '').upper(); g = str(r[2] or '')
    if 'HENDI' in mm or 'ARKTIC' in mm or g.startswith('8711369'): return 'Hendi'
    if 'ARCOS' in mm or g.startswith('84210022'): return 'Arcos'
    if 'CUPERINOX' in mm: return 'Cuperinox'
    if 'VINFER' in mm: return 'Vinfer'
    return mm.title()[:18] if mm.strip() else 'Sin clasificar'

def leer(p):
    rows = list(load_workbook(p, data_only=True, read_only=True)['Export']
                .iter_rows(values_only=True))
    filt = next((str(c) for r in rows for c in r if c and 'Applied filters' in str(c)), '')
    m = re.search(r'after (\S+) and is before (\S+)', filt)
    tot = next((num(r[4]) for r in rows if str(r[0] or '').strip() == 'Total'), None)
    return ([r for r in rows[1:] if r[1]],
            (m.group(1), m.group(2)) if m else ('?', '?'), tot)

Y = {}
for a, (p, meses) in FICHEROS.items():
    rows, rango, tot = leer(p)
    v = sum(num(r[4]) for r in rows); u = sum(num(r[5]) for r in rows)
    Y[a] = dict(rows=rows, meses=meses, desde=rango[0], hasta=rango[1],
                total=round(v, 2), uds=int(u), refs=len(rows),
                mes=round(v / meses, 2), cuadra=(tot is None or abs(tot - v) < 1))

def dim(a, fn):
    D = collections.defaultdict(lambda: dict(v=0., u=0., n=0))
    for r in Y[a]['rows']:
        d = D[fn(r)]; d['v'] += num(r[4]); d['u'] += num(r[5]); d['n'] += 1
    return D

def comparar(fn):
    """Una fila por clave, con los dos anos en EUR/mes."""
    A, B = dim('2025', fn), dim('2026', fn)
    out = []
    for k in sorted(set(A) | set(B), key=lambda k: -B.get(k, dict(v=0))['v']):
        a, b = A.get(k), B.get(k)
        ma = a['v'] / Y['2025']['meses'] if a else 0.0
        mb = b['v'] / Y['2026']['meses'] if b else 0.0
        out.append(dict(k=k,
                        v25=round(a['v'], 2) if a else 0, v26=round(b['v'], 2) if b else 0,
                        u25=int(a['u']) if a else 0, u26=int(b['u']) if b else 0,
                        n25=a['n'] if a else 0, n26=b['n'] if b else 0,
                        m25=round(ma, 2), m26=round(mb, 2),
                        var=round(100 * (mb / ma - 1), 1) if ma else None))
    return out

# top referencias de 2026, con su 2025 al lado
A26 = dim('2026', lambda r: str(r[9] or r[2] or ''))
A25 = dim('2025', lambda r: str(r[9] or r[2] or ''))
META = {}
for a in ('2025', '2026'):
    for r in Y[a]['rows']:
        k = str(r[9] or r[2] or '')
        META.setdefault(k, dict(nombre=str(r[3] or '')[:70], plat=PLAT.get(r[1], str(r[1])),
                                marca=marca(r), cat=str(r[7] or '')))
top = []
for k, d in sorted(A26.items(), key=lambda kv: -kv[1]['v'])[:70]:
    a = A25.get(k)
    top.append(dict(mpn=k, **META[k], v26=round(d['v'], 2), u26=int(d['u']),
                    v25=round(a['v'], 2) if a else 0, u25=int(a['u']) if a else 0,
                    pvp=round(d['v'] / d['u'], 2) if d['u'] else None,
                    nuevo=(a is None)))

json.dump(dict(
    anios=['2025', '2026'],
    meta={a: {k: v for k, v in Y[a].items() if k != 'rows'} for a in Y},
    plataformas=[{**g, 'k': PLAT.get(g['k'], g['k'])} for g in comparar(lambda r: r[1])],
    marcas=comparar(marca),
    categorias=comparar(lambda r: str(r[6] or 'Sin categoría')),
    top=top), open('makro.json', 'w'), ensure_ascii=False, indent=1)

a, b = Y['2025'], Y['2026']
print(f"2025 {a['desde']}→{a['hasta']}: {a['total']:>11,.0f} € · {a['uds']:>6,.0f} uds · "
      f"{a['refs']:>5d} refs · {a['mes']:>8,.0f} €/mes · cuadra {a['cuadra']}")
print(f"2026 {b['desde']}→{b['hasta']}: {b['total']:>11,.0f} € · {b['uds']:>6,.0f} uds · "
      f"{b['refs']:>5d} refs · {b['mes']:>8,.0f} €/mes · cuadra {b['cuadra']}")
print(f"\nCANAL: {b['mes']/a['mes']:.2f}× en €/mes  ({(b['mes']/a['mes']-1)*100:+.0f}%)")
for lab, g in (('PLATAFORMA', 'plataformas'), ('MARCA', 'marcas')):
    J = json.load(open('makro.json'))[g]
    print(f"\n{lab:20s}{'€/mes 25':>10s}{'€/mes 26':>10s}{'var':>9s}{'refs 25':>9s}{'refs 26':>9s}")
    for x in J:
        var = f"{x['var']:+8.0f}%" if x['var'] is not None else '   nuevo'
        print(f"{x['k'][:20]:20s}{x['m25']:10,.0f}{x['m26']:10,.0f}{var}"
              f"{x['n25']:9d}{x['n26']:9d}")
