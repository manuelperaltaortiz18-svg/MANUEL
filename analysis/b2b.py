"""Canal B2B por comercial. Lee la hoja MENSUAL, que es la fiable:
la hoja ANUAL tiene la columna de 2026 desplazada una fila desde ANGEL."""
import json, collections
from openpyxl import load_workbook

XLS = 'b2b_raw/comerciales.xlsx'
MESES = ['ene','feb','mar','abr','may','jun','jul','ago','sep','oct','nov','dic']

def leer():
    ws = load_workbook(XLS, data_only=True, read_only=True)['MENSUAL COMERCIALES']
    D = collections.defaultdict(dict)       # {comercial: {anio: [12 meses]}}
    anio = None
    for r in ws.iter_rows(values_only=True):
        a = r[0]
        if a is None: continue
        a = str(a).strip()
        if a.isdigit() and len(a) == 4:
            anio = int(a); continue
        if anio and a and a != 'COMERCIAL':
            D[a][anio] = [(x if isinstance(x, (int, float)) else 0.0) for x in r[1:13]]
    return D

D = leer()
ANIOS = sorted({y for v in D.values() for y in v})

def cierre(y):
    """Ultimo mes con dato de ese anio en el conjunto del canal."""
    m = 0
    for v in D.values():
        for i, x in enumerate(v.get(y, []), 1):
            if x: m = max(m, i)
    return m

CIERRE = {y: cierre(y) for y in ANIOS}
ULT = max(ANIOS); W = CIERRE[ULT]           # ventana comparable: ene..W

filas = []
for c, v in sorted(D.items(), key=lambda kv: -sum(kv[1].get(ULT, [0]))):
    anual = {y: round(sum(v.get(y, [])), 2) for y in ANIOS}
    lfl = {y: round(sum(v.get(y, [0]*12)[:W]), 2) for y in ANIOS}
    act = [i+1 for i, x in enumerate(v.get(ULT, [])) if x]
    filas.append(dict(comercial=c, anual=anual, lfl=lfl,
                      meses={str(y): v.get(y, [0]*12) for y in ANIOS},
                      activo=bool(act), ultimo_mes=max(act) if act else None,
                      alta=min([y for y in ANIOS if anual[y] > 0], default=None),
                      baja=max([y for y in ANIOS if anual[y] > 0], default=None)))

# Relevos: un comercial que hereda la plaza de otro. Comparar a la persona
# oculta que la plaza siguio existiendo; hay que sumarlos.
RELEVOS = [dict(plaza='Madrid', de='ÁNGEL', a='MANUEL')]
relevos = []
for r in RELEVOS:
    fa = next((f for f in filas if f['comercial'] == r['de']), None)
    fb = next((f for f in filas if f['comercial'] == r['a']), None)
    if not (fa and fb): continue
    relevos.append(dict(**r,
        lfl={y: round(fa['lfl'][y] + fb['lfl'][y], 2) for y in ANIOS},
        anual={y: round(fa['anual'][y] + fb['anual'][y], 2) for y in ANIOS}))

TOT = {y: round(sum(f['anual'][y] for f in filas), 2) for y in ANIOS}
LFL = {y: round(sum(f['lfl'][y] for f in filas), 2) for y in ANIOS}
MES = {str(y): [round(sum(f['meses'][str(y)][i] for f in filas), 2) for i in range(12)] for y in ANIOS}

# canal sin las plazas en relevo: aisla la caida del resto de la red
EXCL = {n for r in RELEVOS for n in (r['de'], r['a'])}
RESTO = {y: round(sum(f['lfl'][y] for f in filas if f['comercial'] not in EXCL), 2) for y in ANIOS}

json.dump(dict(anios=ANIOS, cierre=CIERRE, ventana=W, meses_lab=MESES,
               total=TOT, lfl=LFL, mensual=MES, filas=filas,
               relevos=relevos, resto=RESTO),
          open('b2b.json', 'w'), ensure_ascii=False, indent=1)

print(f"{'comercial':14s}" + ''.join(f"{y:>11d}" for y in ANIOS))
for f in filas:
    print(f"{f['comercial'][:14]:14s}" + ''.join(f"{f['anual'][y]:11,.0f}" for y in ANIOS))
print(f"{'TOTAL':14s}" + ''.join(f"{TOT[y]:11,.0f}" for y in ANIOS))
print(f"{'cierre (mes)':14s}" + ''.join(f"{CIERRE[y]:11d}" for y in ANIOS))
print(f"\nVENTANA COMPARABLE ene-{MESES[W-1]} ({W} meses)")
print(f"{'comercial':14s}" + ''.join(f"{y:>11d}" for y in ANIOS) + f"{'var ult':>10s}")
for f in filas:
    a, b = f['lfl'][ANIOS[-2]], f['lfl'][ULT]
    var = f"{(b/a-1)*100:+9.0f}%" if a else ("  nuevo" if b else "       —")
    print(f"{f['comercial'][:14]:14s}" + ''.join(f"{f['lfl'][y]:11,.0f}" for y in ANIOS) + var)
a, b = LFL[ANIOS[-2]], LFL[ULT]
print(f"{'TOTAL':14s}" + ''.join(f"{LFL[y]:11,.0f}" for y in ANIOS) + f"{(b/a-1)*100:+9.1f}%")

print("\nPLAZAS EN RELEVO (ventana comparable)")
for r in relevos:
    print(f"  {r['plaza']:10s} ({r['de']} -> {r['a']})" + ''.join(f"{r['lfl'][y]:11,.0f}" for y in ANIOS))
print("\nRESTO DE LA RED, sin las plazas en relevo (ventana comparable)")
print(f"  {'resto':10s}" + ''.join(f"{RESTO[y]:11,.0f}" for y in ANIOS)
      + f"   var ultimo: {(RESTO[ULT]/RESTO[ANIOS[-2]]-1)*100:+.1f}%")
